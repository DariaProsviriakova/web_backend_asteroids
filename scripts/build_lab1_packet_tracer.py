from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Inches, Pt


OUT_DIR = Path("output/lab1_packet_tracer")
REPORT_DOCX = OUT_DIR / "lab1_packet_tracer_report.docx"
REPORT_MD = OUT_DIR / "lab1_packet_tracer_report.md"
COMMANDS_TXT = OUT_DIR / "lab1_ios_commands.txt"
INSTALL_MD = OUT_DIR / "packet_tracer_macos_install.md"


COMMANDS = """Cisco IOS commands for Lab 1
Replace X, Y, A, B, K before entering commands:
X - group number, Y - variant number, A - number of PCs in segment 1,
B - number of PCs in segment 2, K - allowed number of MAC addresses from
segment 1 on the trunk, with K < A for the Port-Security test.

Example IP plan:
Network: 196.X.Y.0/24
SW1 management VLAN 10: 196.X.Y.2
SW2 management VLAN 10: 196.X.Y.3
PC_Ai: 196.X.Y.(10+i), mask 255.255.255.0
PC_Bj: 196.X.Y.(100+j), mask 255.255.255.0
Gateway: leave blank, because all hosts are in one subnet and no router is used.

SW1
---
enable
configure terminal
hostname SW1
no ip domain-lookup
enable secret class
service password-encryption

vlan 10
 name USERS
exit
vlan 99
 name NATIVE_UNUSED
exit

interface range fa0/1 - <A>
 switchport mode access
 switchport access vlan 10
 spanning-tree portfast
 no shutdown
exit

interface fa0/24
 description TRUNK_TO_SW2
 switchport mode trunk
 switchport trunk allowed vlan 10,99
 switchport trunk native vlan 99
 no shutdown
exit

interface vlan 10
 ip address 196.X.Y.2 255.255.255.0
 no shutdown
exit

line console 0
 password console123
 login
 logging synchronous
exit

line vty 0 4
 password telnet123
 login
 transport input telnet
exit

end
copy running-config startup-config


SW2
---
enable
configure terminal
hostname SW2
no ip domain-lookup
enable secret class
service password-encryption

vlan 10
 name USERS
exit
vlan 99
 name NATIVE_UNUSED
exit

interface range fa0/1 - <B>
 switchport mode access
 switchport access vlan 10
 spanning-tree portfast
 no shutdown
exit

interface fa0/24
 description TRUNK_TO_SW1
 switchport mode trunk
 switchport trunk allowed vlan 10,99
 switchport trunk native vlan 99
 switchport port-security
 switchport port-security maximum <K>
 switchport port-security violation restrict
 switchport port-security mac-address sticky
 no shutdown
exit

interface vlan 10
 ip address 196.X.Y.3 255.255.255.0
 no shutdown
exit

line console 0
 password console123
 login
 logging synchronous
exit

line vty 0 4
 password telnet123
 login
 transport input telnet
exit

end
copy running-config startup-config


Verification commands
---------------------
show vlan brief
show interfaces trunk
show port-security interface fa0/24
show port-security address
show mac address-table
ping 196.X.Y.101
telnet 196.X.Y.2

If Packet Tracer rejects Port-Security on a trunk port, keep the trunk settings
and demonstrate the same logic on the access port of the limited segment:

interface fa0/1
 switchport mode access
 switchport access vlan 10
 switchport port-security
 switchport port-security maximum 1
 switchport port-security violation restrict
 switchport port-security mac-address sticky
"""


REPORT_MD_TEXT = """# Лабораторная работа 1 Локальная сеть на коммутаторах Cisco Packet Tracer

## Цель работы

Закрепить навыки проектирования локальной вычислительной сети в Cisco Packet
Tracer, настроить коммутаторы второго уровня, проверить передачу пакетов между
сегментами сети, сравнить кадры на access и trunk участках, а также
продемонстрировать ограничение доступа с помощью Port-Security.

## Исходные данные

В методических указаниях параметры задаются переменными:

| Параметр | Значение |
| --- | --- |
| X | номер группы |
| Y | номер варианта |
| A | количество компьютеров в первом сегменте |
| B | количество компьютеров во втором сегменте |
| K | разрешенное число MAC-адресов для проверки Port-Security, выбирается K < A |

Все устройства находятся в одной подсети `196.X.Y.0/24`. Для пользовательских
портов используется VLAN 10, а на trunk-линии native VLAN изменен на VLAN 99.
Так кадры VLAN 10 на участке между коммутаторами передаются с тегом IEEE
802.1Q, а кадры между ПК и коммутатором остаются нетегированными.

## Топология

```text
PC_A1 ... PC_AA -- Fa0/1..Fa0/A   SW1   Fa0/24 == trunk == Fa0/24   SW2   Fa0/1..Fa0/B -- PC_B1 ... PC_BB
                                      ^
                                      |
                         Laptop-Console -- RS-232 console
```

Используются два коммутатора Cisco 2960. Рабочие станции соединяются с
коммутаторами медным прямым кабелем. Коммутаторы соединяются медным
кроссоверным кабелем через FastEthernet0/24. Один из коммутаторов настраивается
через консольный RS-232 порт.

## Адресный план

| Устройство | Интерфейс | IP адрес | Маска |
| --- | --- | --- | --- |
| SW1 | VLAN 10 | 196.X.Y.2 | 255.255.255.0 |
| SW2 | VLAN 10 | 196.X.Y.3 | 255.255.255.0 |
| PC_A1 | FastEthernet0 | 196.X.Y.11 | 255.255.255.0 |
| PC_Ai | FastEthernet0 | 196.X.Y.(10+i) | 255.255.255.0 |
| PC_B1 | FastEthernet0 | 196.X.Y.101 | 255.255.255.0 |
| PC_Bj | FastEthernet0 | 196.X.Y.(100+j) | 255.255.255.0 |

Шлюз по умолчанию для ПК не задается, потому что все узлы находятся в одной
подсети и маршрутизация между сетями не используется.

## Настройка коммутаторов

Полный набор команд находится в файле `lab1_ios_commands.txt`. Перед вводом
нужно заменить `X`, `Y`, `<A>`, `<B>` и `<K>` на значения варианта.

Для SW1 создаются VLAN 10 и VLAN 99, пользовательские порты переводятся в
access VLAN 10, порт Fa0/24 настраивается как trunk с native VLAN 99. Также
задается management IP адрес 196.X.Y.2 и пароль для консоли.

Для SW2 повторяется базовая VLAN-настройка. На trunk-порту Fa0/24 дополнительно
включается Port-Security: максимальное число разрешенных MAC-адресов задается
равным K, режим нарушения установлен `restrict`, а адреса изучаются как sticky.

## Проверка связи

После настройки выполняются команды:

```text
show vlan brief
show interfaces trunk
show port-security interface fa0/24
show port-security address
show mac address-table
```

Связь проверяется ICMP-запросами между ПК разных сегментов, например от PC_A1
к PC_B1:

```text
ping 196.X.Y.101
```

При корректной настройке VLAN и trunk-линии ответы на ping должны приходить со
статистикой 100 процентов успешных эхо-ответов для разрешенных узлов.

## Сравнение кадров access и trunk

В режиме Simulation нужно включить фильтр ICMP и отправить Simple PDU от PC_A1
к PC_B1. На участке PC_A1 - SW1 кадр передается как обычный Ethernet-кадр без
VLAN-тега, потому что порт работает в access-режиме. На участке SW1 - SW2 кадр
проходит через trunk и содержит IEEE 802.1Q tag с идентификатором VLAN 10.

Native VLAN на trunk-портах задан как VLAN 99. Поэтому пользовательский трафик
VLAN 10 не считается native-трафиком и отображается в Packet Tracer как
тегированный. Native VLAN должна совпадать на обеих сторонах trunk-соединения.

## Демонстрация Port-Security

Для проверки выбирается значение K меньше количества компьютеров в первом
сегменте. Сначала K разных ПК из первого сегмента отправляют ping на PC_B1.
Коммутатор SW2 изучает их MAC-адреса как sticky и пропускает трафик. Затем
следующий ПК из первого сегмента отправляет ping на PC_B1. Так как лимит
MAC-адресов превышен, SW2 применяет режим `restrict`: кадры нарушителя
отбрасываются, счетчик нарушений увеличивается, а разрешенные MAC-адреса
остаются рабочими.

Проверка состояния:

```text
show port-security interface fa0/24
show port-security address
```

Если конкретная версия Packet Tracer не разрешает включить Port-Security на
trunk-порту, для демонстрации можно применить те же команды на access-порту
ограничиваемого сегмента и указать это как особенность симулятора.

## Настройка удаленного доступа

Дополнительно включен доступ Telnet к коммутатору. Для проверки с любого ПК в
той же подсети запускается:

```text
telnet 196.X.Y.2
```

Пароль VTY: `telnet123`. Для входа в привилегированный режим используется
пароль `class`.

## Ответы на контрольные вопросы

1. В работе использовались физические технологии Ethernet по медной витой паре
для подключения ПК и коммутаторов, а также RS-232 для консольного подключения к
коммутатору. Для связи ПК с коммутатором применяется прямой медный кабель, для
связи коммутаторов по условию работы - кроссоверный кабель.

2. Access-порт принадлежит одному VLAN и передает кадры без VLAN-тегов. Такой
режим применяется для конечных устройств. Trunk-порт переносит трафик
нескольких VLAN между сетевыми устройствами и добавляет к кадрам тег IEEE
802.1Q, кроме кадров native VLAN, которые передаются без тега.

3. Port-Security нужен для ограничения количества или списка MAC-адресов,
которым разрешено работать через порт коммутатора. Это снижает риск
несанкционированного подключения устройств, подмены рабочих станций и атак с
переполнением таблицы MAC-адресов.

4. На канальном уровне использовались Ethernet IEEE 802.3 для передачи кадров
в локальной сети и IEEE 802.1Q для маркировки VLAN на trunk-соединении. Также
коммутатор может использовать служебные протоколы уровня 2, например STP,
чтобы предотвращать петли. На физическом уровне применялись медная витая пара
Ethernet и RS-232 для консоли. ICMP и IP использовались для проверки связи, но
они относятся к более высоким уровням модели OSI.

5. Основные режимы нарушения Port-Security: `protect`, `restrict` и `shutdown`.
В режиме `protect` кадры от неизвестных MAC-адресов отбрасываются без
уведомлений. В режиме `restrict` такие кадры отбрасываются, а счетчик нарушений
увеличивается и формируются сообщения. В режиме `shutdown`, который является
режимом по умолчанию, порт переводится в error-disabled состояние и перестает
передавать трафик до восстановления администратором.

6. VLAN - это логическая локальная сеть, то есть отдельный широковещательный
домен внутри одной физической коммутируемой инфраструктуры. VLAN применяют для
изоляции групп устройств, уменьшения широковещательного трафика, повышения
безопасности и удобного разделения сети на отделы или сервисы. Для обмена между
разными VLAN требуется маршрутизация.

## Вывод

В ходе лабораторной работы построена локальная сеть из двух коммутируемых
сегментов в одной подсети. Настроены access-порты для конечных устройств,
trunk-соединение между коммутаторами с измененной native VLAN, проверена
передача ICMP-пакетов, показано отличие тегированных кадров trunk от
нетегированных кадров access-сегмента, а также продемонстрировано ограничение
числа MAC-адресов с помощью Port-Security.
"""


INSTALL_MD_TEXT = """# Cisco Packet Tracer на macOS

Официальная страница загрузки: <https://www.netacad.com/resources/lab-downloads?courseLang=en-US>

На 15 сентября 2026 года страница Cisco NetAcad показывает пакет
`Packet Tracer 9.0.1 MacOS 64bit`. Системные требования на странице:
macOS 12 или новее, amd64/x86-64 CPU, 4 GB свободной RAM и 1.4 GB свободного
места на диске. Для скачивания требуется вход в NetAcad.

Официальная инструкция Cisco:
<https://www.netacad.com/skillsforall/files/Cisco_Packet_Tracer_Download_and_Installation_Instructions.pdf>

Порядок установки:

1. Открыть NetAcad Resource Hub.
2. Войти в Cisco NetAcad или Skills For All.
3. Нажать `Packet Tracer 9.0.1 MacOS 64bit`.
4. Сохранить `.dmg` файл в Downloads.
5. Открыть `.dmg`, выбрать `Open`, затем `Continue`.
6. Принять лицензионное соглашение и нажать `Install`.
7. При запросе macOS ввести пароль или PIN учетной записи.
8. После установки открыть Finder, перейти в Applications, открыть папку Cisco
   Packet Tracer и запустить приложение.
9. При первом запуске нажать `LOGIN` и авторизоваться через NetAcad.

Примечание: прямую ссылку на `.dmg` Cisco не показывает без входа в аккаунт.
Не стоит скачивать Packet Tracer с зеркал, потому что это учебный сетевой
симулятор с закрытым установщиком, и безопаснее брать его только из NetAcad.
"""


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    p.paragraph_format.space_after = Pt(2)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_text(cell, header, bold=True)
        set_cell_shading(cell, "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    doc.add_paragraph()
    return table


def add_code_block(doc, text):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(line)
        run.font.name = "Courier New"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")
        run.font.size = Pt(9)
    doc.add_paragraph()


def add_heading(doc, text, level):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    return paragraph


def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        rest = p.add_run(text[len(bold_lead):])
    else:
        rest = p.add_run(text)
    for run in p.runs:
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        run.font.size = Pt(11)
    return p


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    styles["Normal"].font.size = Pt(11)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
        p_pr = style._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Лабораторная работа 1 Локальная сеть на коммутаторах Cisco Packet Tracer")
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(18)
    run.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    p_bdr = title._p.get_or_add_pPr().find(qn("w:pBdr"))
    if p_bdr is not None:
        title._p.get_or_add_pPr().remove(p_bdr)

    add_paragraph(
        doc,
        "Отчет описывает построение двухсегментной локальной сети в Cisco Packet Tracer, "
        "настройку access и trunk портов, изменение native VLAN, проверку передачи ICMP-пакетов "
        "и демонстрацию ограничения MAC-адресов с помощью Port-Security.",
    )

    add_heading(doc, "Цель работы", 1)
    add_paragraph(
        doc,
        "Закрепить навыки проектирования локальной вычислительной сети, настройки коммутаторов "
        "в Cisco IOS, анализа кадров в режиме Simulation и проверки защитных настроек порта.",
    )

    add_heading(doc, "Исходные данные", 1)
    add_table(
        doc,
        ["Параметр", "Значение"],
        [
            ["X", "номер группы"],
            ["Y", "номер варианта"],
            ["A", "количество ПК в первом сегменте"],
            ["B", "количество ПК во втором сегменте"],
            ["K", "лимит MAC-адресов для проверки Port-Security, K меньше A"],
        ],
    )
    add_paragraph(
        doc,
        "Все устройства размещаются в подсети 196.X.Y.0/24. Пользовательские порты работают "
        "в VLAN 10, а native VLAN на trunk-соединении изменена на VLAN 99. Поэтому трафик VLAN 10 "
        "между коммутаторами передается с тегом IEEE 802.1Q, а трафик на участке ПК - коммутатор "
        "остается нетегированным.",
    )

    add_heading(doc, "Топология", 1)
    add_code_block(
        doc,
        """PC_A1 ... PC_AA -- Fa0/1..Fa0/A   SW1   Fa0/24 == trunk == Fa0/24   SW2   Fa0/1..Fa0/B -- PC_B1 ... PC_BB
                                      ^
                                      |
                         Laptop-Console -- RS-232 console""",
    )
    add_paragraph(
        doc,
        "Рабочие станции подключаются к коммутаторам медным прямым кабелем. Коммутаторы соединяются "
        "между собой кроссоверным кабелем. Один из коммутаторов настраивается через консольный RS-232 порт.",
    )

    add_heading(doc, "Адресный план", 1)
    add_table(
        doc,
        ["Устройство", "Интерфейс", "IP адрес", "Маска"],
        [
            ["SW1", "VLAN 10", "196.X.Y.2", "255.255.255.0"],
            ["SW2", "VLAN 10", "196.X.Y.3", "255.255.255.0"],
            ["PC_A1", "FastEthernet0", "196.X.Y.11", "255.255.255.0"],
            ["PC_Ai", "FastEthernet0", "196.X.Y.(10+i)", "255.255.255.0"],
            ["PC_B1", "FastEthernet0", "196.X.Y.101", "255.255.255.0"],
            ["PC_Bj", "FastEthernet0", "196.X.Y.(100+j)", "255.255.255.0"],
        ],
    )
    add_paragraph(
        doc,
        "Шлюз по умолчанию на ПК не задается, так как все узлы находятся в одной подсети и маршрутизатор "
        "в топологии не используется.",
    )

    add_heading(doc, "Основная конфигурация", 1)
    add_paragraph(
        doc,
        "На обоих коммутаторах создаются VLAN 10 для пользовательского трафика и VLAN 99 как native VLAN "
        "для trunk-соединения. На SW2 на trunk-порту дополнительно включается Port-Security с режимом "
        "нарушения restrict и sticky-изучением MAC-адресов.",
    )
    add_code_block(
        doc,
        """interface fa0/24
 switchport mode trunk
 switchport trunk allowed vlan 10,99
 switchport trunk native vlan 99
 switchport port-security
 switchport port-security maximum <K>
 switchport port-security violation restrict
 switchport port-security mac-address sticky""",
    )

    add_heading(doc, "Проверка", 1)
    add_paragraph(
        doc,
        "Корректность настройки проверяется командами show vlan brief, show interfaces trunk, "
        "show port-security interface fa0/24, show port-security address и ping между ПК разных сегментов.",
    )
    add_code_block(
        doc,
        """show vlan brief
show interfaces trunk
show port-security interface fa0/24
show port-security address
ping 196.X.Y.101""",
    )
    add_paragraph(
        doc,
        "Для разрешенных MAC-адресов ICMP-запросы должны проходить успешно. После превышения лимита K "
        "кадры от нового MAC-адреса отбрасываются, а счетчик нарушений Port-Security увеличивается.",
    )

    add_heading(doc, "Сравнение кадров", 1)
    add_paragraph(
        doc,
        "В режиме Simulation кадр между ПК и коммутатором отображается как обычный Ethernet-кадр без "
        "VLAN-тега. На участке между SW1 и SW2 этот же трафик проходит через trunk и содержит тег "
        "IEEE 802.1Q с VLAN ID 10. Native VLAN 99 настроена одинаково на обеих сторонах trunk.",
    )

    add_heading(doc, "Удаленный доступ", 1)
    add_paragraph(
        doc,
        "Дополнительно настроен Telnet-доступ к SW1 по адресу 196.X.Y.2. Пароль VTY: telnet123. "
        "Для входа в привилегированный режим используется enable secret class.",
    )

    add_heading(doc, "Ответы на контрольные вопросы", 1)
    questions = [
        (
            "Какие протоколы физического уровня были использованы в данной работе",
            "Ethernet по медной витой паре для подключений ПК и коммутаторов, а также RS-232 для "
            "консольного подключения. ПК соединяются с коммутаторами прямым кабелем, коммутаторы по "
            "условию работы - кроссоверным кабелем.",
        ),
        (
            "Чем отличаются режимы trunk и access",
            "Access-порт относится к одному VLAN и передает кадры без тегов. Trunk-порт переносит "
            "несколько VLAN между сетевыми устройствами и маркирует кадры тегом IEEE 802.1Q, кроме "
            "кадров native VLAN.",
        ),
        (
            "Для чего необходимо использовать Port-Security",
            "Port-Security ограничивает количество или конкретный список MAC-адресов на порту, "
            "уменьшая риск несанкционированного подключения и переполнения MAC-таблицы.",
        ),
        (
            "Какие протоколы канального и физического уровня были использованы",
            "На канальном уровне использовались Ethernet IEEE 802.3 и IEEE 802.1Q для VLAN-тегов. "
            "На физическом уровне применялись медная витая пара Ethernet и RS-232. ICMP и IP применялись "
            "для проверки связи, но относятся к более высоким уровням.",
        ),
        (
            "Какие режимы работы Port-Security существуют",
            "Protect отбрасывает кадры неизвестных MAC-адресов без уведомлений. Restrict отбрасывает "
            "такие кадры и увеличивает счетчик нарушений. Shutdown переводит порт в error-disabled "
            "состояние и блокирует передачу до восстановления администратором.",
        ),
        (
            "Что такое VLAN и для чего он применяется",
            "VLAN - логическая локальная сеть и отдельный широковещательный домен внутри коммутируемой "
            "инфраструктуры. VLAN применяют для изоляции устройств, снижения широковещательного трафика, "
            "повышения безопасности и удобного разделения сети.",
        ),
    ]
    for idx, (question, answer) in enumerate(questions, start=1):
        add_paragraph(doc, f"{idx}. {question}? {answer}")

    add_heading(doc, "Вывод", 1)
    add_paragraph(
        doc,
        "В лабораторной работе построена сеть из двух коммутируемых сегментов в одной подсети, "
        "настроены access-порты, trunk-соединение с измененной native VLAN, проверена передача ICMP "
        "между сегментами и продемонстрировано ограничение MAC-адресов с помощью Port-Security.",
    )

    doc.save(REPORT_DOCX)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_MD.write_text(REPORT_MD_TEXT, encoding="utf-8")
    COMMANDS_TXT.write_text(COMMANDS, encoding="utf-8")
    INSTALL_MD.write_text(INSTALL_MD_TEXT, encoding="utf-8")
    build_docx()


if __name__ == "__main__":
    main()
