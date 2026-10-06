from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/daraprosvirakova/web_frontend_asteroids")
SCREENSHOT_DIR = ROOT / "output" / "lr3_real_screenshots"
OUTPUT_PATH = ROOT / "output" / "Отчет ЛР3 Asteroid Dates Просвиракова реальные скрины.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table, header_fill: str = "1F2937") -> None:
    table.autofit = True
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_borders(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F3F6FA")


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Параметр"
    table.rows[0].cells[1].text = "Значение"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    style_table(table)


def add_methods_table(doc: Document) -> None:
    rows = [
        ("GET", "/api/dates?maxMonth=7", "Список опубликованных услуг с фильтром по месяцу."),
        ("POST", "/api/dates", "Создание черновика. multipart/form-data: designation, image, video."),
        ("GET", "/api/dates/draft", "Получение текущего черновика пользователя без передачи id."),
        ("PUT", "/api/dates/publish", "Публикация текущего черновика. Статус меняется только backend-логикой."),
        ("GET", "/api/dates/feed", "Первая запись ленты опубликованных услуг."),
        ("GET", "/api/dates/feed?id=2&next=true", "Следующая запись ленты через запрос к БД, id может иметь пропуски."),
        ("POST", "/api/dates/2/like", "like=1 ставит лайк, like=0 отменяет. Количество в ответе ограничено 5."),
        ("DELETE", "/api/dates/:id", "Soft delete только своей услуги."),
        ("POST", "/api/users/register", "Регистрация пользователя."),
        ("POST", "/api/users/auth, /api/users/logout", "Заглушки аутентификации и деавторизации для следующей лабораторной."),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ("Метод", "URL", "Назначение")
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    style_table(table)


def add_tables_table(doc: Document) -> None:
    rows = [
        (
            "asteroid_dates",
            "id, designation, short_description, status, image_url, video_url, created_at, formed_at, creator_id, approach_month, approach_day, minimum_distance_au",
            "Услуги/даты сближения. Статусы: draft, published, deleted.",
        ),
        (
            "observers",
            "id, full_name, email",
            "Пользователи системы.",
        ),
        (
            "observer_date_likes",
            "id, observer_id, asteroid_date_id",
            "Связь пользователя и понравившейся услуги; уникальность observer_id + asteroid_date_id.",
        ),
    ]
    table = doc.add_table(rows=1, cols=3)
    headers = ("Таблица", "Поля", "Назначение")
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    style_table(table)


def add_screenshot(doc: Document, path: str, caption: str, width: float = 6.55) -> None:
    image_path = SCREENSHOT_DIR / path
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(caption).bold = True
    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.add_run().add_picture(str(image_path), width=Inches(width))


def set_document_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)


def build_report() -> None:
    doc = Document()
    set_document_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Отчет по лабораторной работе 3").bold = True
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Веб сервис Asteroid Dates для SPA").bold = True

    add_kv_table(
        doc,
        [
            ("Студент", "Дарья Просвиракова"),
            ("Проект", "Asteroid Dates"),
            ("Backend", "NestJS, TypeORM, PostgreSQL"),
            ("Текущий пользователь в лабораторной", "id=1, Дарья Просвиракова, daria@example.com"),
            ("Дата проверки", "6 октября 2026"),
        ],
    )

    doc.add_paragraph(
        "Цель работы: создать REST веб сервис на backend для использования из SPA, подключить его к базе данных, "
        "реализовать итоговую бизнес логику без полноценной авторизации и проверить методы через Postman."
    )
    doc.add_paragraph(
        "В лабораторной работе реализованы домены услуг и пользователей. Системные поля не передаются клиентом для изменения: "
        "id, статусы, пользователь, даты и вычисляемые признаки задаются backend-логикой."
    )

    add_heading(doc, "HTTP методы", 1)
    add_methods_table(doc)

    add_heading(doc, "Таблицы базы данных", 1)
    add_tables_table(doc)

    doc.add_page_break()
    add_heading(doc, "Демонстрация запросов Postman", 1)
    postman_shots = [
        ("01_postman_get_dates_filter.jpg", "Скриншот 1. GET /api/dates?maxMonth=7, список опубликованных услуг с фильтром."),
        ("02_postman_post_date_files.jpg", "Скриншот 2. POST /api/dates, добавление черновика с image и video как файлами."),
        ("03_postman_get_draft.jpg", "Скриншот 3. GET /api/dates/draft, получение текущего черновика."),
        ("04_postman_put_publish.jpg", "Скриншот 4. PUT /api/dates/publish, публикация черновика."),
        ("05_postman_get_feed_first.jpg", "Скриншот 5. GET /api/dates/feed, лента без id."),
        ("06_postman_get_feed_next.jpg", "Скриншот 6. GET /api/dates/feed?id=2&next=true, следующая запись ленты."),
        ("07_postman_like_set.jpg", "Скриншот 7. POST /api/dates/2/like, установка лайка like=1."),
        ("08_postman_like_unset.jpg", "Скриншот 8. POST /api/dates/2/like, отмена лайка like=0."),
        ("09_postman_delete_soft.jpg", "Скриншот 9. DELETE /api/dates/21, soft delete своей услуги."),
        ("10_postman_register_user.jpg", "Скриншот 10. POST /api/users/register, регистрация нового пользователя."),
    ]
    for index, (filename, caption) in enumerate(postman_shots, 1):
        add_screenshot(doc, filename, caption)
        if index in (2, 4, 6, 8, 10):
            doc.add_page_break()

    add_heading(doc, "Проверка данных через SELECT", 1)
    select_shots = [
        ("11_vscode_select_dates.jpg", "Скриншот 11. SELECT по asteroid_dates: созданные услуги переведены в deleted через soft delete."),
        ("12_vscode_select_likes.jpg", "Скриншот 12. SELECT по observer_date_likes: данные лайков после проверки методов."),
        ("13_vscode_select_users.jpg", "Скриншот 13. SELECT по observers: фиксированный пользователь и зарегистрированный пользователь отчета."),
    ]
    for filename, caption in select_shots:
        add_screenshot(doc, filename, caption)
    doc.add_page_break()

    add_heading(doc, "Модели и сериализаторы", 1)
    doc.add_paragraph(
        "Модель AsteroidDate описывает таблицу asteroid_dates, статусы, связи с пользователем и лайками. "
        "Сериализатор формирует ответ API и добавляет признаки isCreator и isLiked для текущего пользователя."
    )
    add_screenshot(doc, "14_vscode_models.jpg", "Скриншот 14. Модель AsteroidDate в TypeORM.", width=6.55)
    add_screenshot(doc, "15_vscode_serializer.jpg", "Скриншот 15. Сериализатор ответа услуги.", width=6.55)
    doc.add_page_break()

    add_heading(doc, "Singleton пользователя и README", 1)
    doc.add_paragraph(
        "Авторизация в лабораторной не реализуется, поэтому текущий создатель зафиксирован константой в singleton-функции. "
        "Сервисные методы получают id пользователя через эту функцию и не принимают creator_id с клиента."
    )
    add_screenshot(doc, "16_vscode_singleton.jpg", "Скриншот 16. Singleton текущего пользователя getCurrentObserver().")
    add_screenshot(doc, "16c_vscode_singleton_usage_methods.jpg", "Скриншот 16.1. Использование CURRENT_OBSERVER_ID в методах сервиса и ограничение лайков до 5.")
    add_screenshot(doc, "17_vscode_readme.jpg", "Скриншот 17. README с запуском, HTTP методами и описанием таблиц.")

    add_heading(doc, "Контрольные вопросы", 1)
    questions = [
        ("Веб сервис", "Программный интерфейс, доступный по сети. В работе сервис отдает и изменяет данные SPA через /api."),
        ("REST", "Архитектурный стиль, где ресурсы представлены URL, а действия выражаются стандартными HTTP методами."),
        ("RPC", "Подход, в котором клиент вызывает удаленную процедуру как функцию, например publishDate или setLike."),
        ("HTTP заголовки и методы", "Методы GET, POST, PUT, DELETE задают действие, а заголовки описывают формат, авторизацию, кэширование и служебные параметры запроса."),
        ("Версии HTTP", "HTTP/1.0 и HTTP/1.1 используют текстовый протокол поверх TCP; HTTP/2 добавляет мультиплексирование; HTTP/3 работает поверх QUIC."),
        ("HTTPS", "HTTP поверх TLS. Шифрует трафик, проверяет сертификат сервера и защищает данные от подмены в канале."),
        ("OSI ISO", "Семислойная модель: физический, канальный, сетевой, транспортный, сеансовый, представления, прикладной уровни."),
    ]
    for question, answer in questions:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{question}: ").bold = True
        paragraph.add_run(answer)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_report()
